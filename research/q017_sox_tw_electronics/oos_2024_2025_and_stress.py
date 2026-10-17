#!/usr/bin/env python3
import csv, io, json, math, re, time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta
from pathlib import Path
import requests
from openpyxl import load_workbook

ROOT=Path('research/q017_sox_tw_electronics')
OUT=ROOT/'oos_2024_2025'; OUT.mkdir(parents=True,exist_ok=True)
START=date(2024,1,1); END=date(2025,12,31)
TWSE='https://www.twse.com.tw/rwd/zh/afterTrading/MI_INDEX'
SOX='https://indexes.nasdaq.com/Index/ExportHistory/SOX'
UA={'User-Agent':'Mozilla/5.0 HENGCE Q017 OOS'}
START_CAP=2_000_000.0

def num(x):
    try:
        s=str(x).replace(',','').strip()
        if s in ('','--','---'): return None
        return float(s)
    except: return None

def dec(b):
    for enc in ('utf-8-sig','utf-8','cp950','big5'):
        try:return b.decode(enc)
        except: pass
    return b.decode('utf-8',errors='replace')

def parse_tw(text):
    rows=list(csv.reader(io.StringIO(text))); hi=None
    for i,row in enumerate(rows):
        j='|'.join(row)
        if '證券代號' in j and '證券名稱' in j and '收盤價' in j: hi=i; break
    if hi is None:return []
    h=[x.strip().replace('\ufeff','') for x in rows[hi]]; out=[]
    for row in rows[hi+1:]:
        row=row+['']*max(0,len(h)-len(row)); r={h[i]:row[i].strip() for i in range(min(len(h),len(row)))}
        t=r.get('證券代號','').replace('=','').replace('"','').strip(); n=r.get('證券名稱','').strip()
        if t and n and re.match(r'^[0-9A-Za-z]{4,8}$',t): out.append((t,n,r.get('開盤價',''),r.get('收盤價','')))
    return out

def fetch_day(d):
    for a in range(6):
        try:
            r=requests.get(TWSE,params={'date':d.strftime('%Y%m%d'),'type':'13','response':'csv'},headers=UA,timeout=(5,20))
            if r.status_code==200 and r.content:
                rows=parse_tw(dec(r.content))
                if rows:return d,'PASS',rows,''
                return d,'NO_DATA',[],''
            err=f'HTTP {r.status_code}'
        except Exception as e:err=repr(e)
        time.sleep(min(15,1.2*(2**a)))
    return d,'ERROR',[],err

def weekdays():
    d=START; a=[]
    while d<=END:
        if d.weekday()<5:a.append(d)
        d+=timedelta(days=1)
    return a

def fetch_twse():
    panel=[]; qa=[]; dates=weekdays()
    with ThreadPoolExecutor(max_workers=2) as ex:
        fs={ex.submit(fetch_day,d):d for d in dates}
        for k,fut in enumerate(as_completed(fs),1):
            d,s,rows,e=fut.result(); qa.append((d.isoformat(),s,len(rows),e))
            for t,n,o,c in rows: panel.append((d.isoformat(),t,n,o,c))
            if k%50==0: print('TWSE',k,'/',len(dates),flush=True)
    qa.sort(); panel.sort()
    with (OUT/'twse.csv').open('w',newline='',encoding='utf-8') as f:
        w=csv.writer(f);w.writerow(['date','ticker','name','open','close']);w.writerows(panel)
    with (OUT/'qa.csv').open('w',newline='',encoding='utf-8') as f:
        w=csv.writer(f);w.writerow(['date','status','rows','error']);w.writerows(qa)
    return panel,qa

def fetch_sox():
    r=requests.get(SOX,params={'startDate':'2024-01-01T00:00:00.000','endDate':'2025-12-31T00:00:00.000','timeOfDay':'EOD'},headers=UA,timeout=(5,30));r.raise_for_status()
    p=OUT/'sox_raw.bin';p.write_bytes(r.content)
    rows=[]
    # Nasdaq export is xlsx here
    wb=load_workbook(p,read_only=True,data_only=True)
    for ws in wb.worksheets:
        for rr in ws.iter_rows(values_only=True):
            if not rr or len(rr)<2:continue
            ds=rr[0]; v=rr[1]
            if isinstance(ds,datetime): d=ds.date()
            elif isinstance(ds,date): d=ds
            else:
                d=None
                for fmt in ('%m/%d/%Y','%Y-%m-%d'):
                    try:d=datetime.strptime(str(ds).strip(),fmt).date();break
                    except:pass
            try:c=float(v)
            except:continue
            if d and START<=d<=END:rows.append((d.isoformat(),c))
    rows=sorted(dict(rows).items())
    with (OUT/'sox.csv').open('w',newline='',encoding='utf-8') as f:
        w=csv.writer(f);w.writerow(['date','close']);w.writerows(rows)
    return rows

def backtest(panel,sox,cost=0.004,hold=3,th=.03,topN=5):
    by_t=defaultdict(list); by_d=defaultdict(list)
    for d,t,n,o,c in panel:
        o=num(o);c=num(c)
        if o is None or c is None:continue
        x={'date':d,'ticker':t,'name':n,'open':o,'close':c};by_t[t].append(x);by_d[d].append(x)
    for t in by_t:by_t[t].sort(key=lambda x:x['date'])
    sr={}; sd=[d for d,_ in sox]
    for i in range(1,len(sox)):sr[sox[i][0]]=sox[i][1]/sox[i-1][1]-1
    def prev_us(d):
        x=[u for u in sd if u<d];return x[-1] if x else None
    aligned={d:sr.get(prev_us(d)) for d in by_d}
    beta={}; hist=defaultdict(list)
    for t,rs in by_t.items():
        pc=None
        for r in rs:
            if pc and aligned.get(r['date']) is not None:
                tr=r['close']/pc-1; obs=hist[t]
                if len(obs)>=30:
                    z=obs[-60:]; mx=sum(a for a,b in z)/len(z); my=sum(b for a,b in z)/len(z);vv=sum((a-mx)**2 for a,b in z)
                    beta[(r['date'],t)]=sum((a-mx)*(b-my) for a,b in z)/vv if vv else 0
                obs.append((aligned[r['date']],tr))
            pc=r['close']
    idx={t:{r['date']:i for i,r in enumerate(rs)} for t,rs in by_t.items()}
    cap=START_CAP;peak=cap;mdd=0;events=[];busy='0000-00-00'
    for d in sorted(by_d):
        if d<=busy:continue
        shock=aligned.get(d)
        if shock is None or shock>-th:continue
        cs=[]
        for r in by_d[d]:
            b=beta.get((d,r['ticker']))
            if b is not None:cs.append((b,r))
        cs=sorted(cs,key=lambda z:z[0],reverse=True)[:topN]
        rr=[]; exits=[]; names=[]
        for b,r in cs:
            i=idx[r['ticker']][d];j=i+hold-1
            if j>=len(by_t[r['ticker']]):continue
            ex=by_t[r['ticker']][j]; ret=ex['close']/r['open']-1-cost;rr.append(ret);exits.append(ex['date']);names.append(r['ticker'])
        if not rr:continue
        pr=sum(rr)/len(rr);cap*=1+pr;peak=max(peak,cap);mdd=min(mdd,cap/peak-1);busy=max(exits)
        events.append({'entry':d,'exit':busy,'ret':pr,'cap':cap,'shock':shock,'names':'|'.join(names)})
    return {'final':cap,'ret':cap/START_CAP-1,'mdd':mdd,'events':events,'wins':sum(x['ret']>0 for x in events),'n':len(events)}

def stress_2026():
    p=ROOT/'backtest_2m'/'event_trades.csv'
    if not p.exists(): return {}
    rows=list(csv.DictReader(p.open(encoding='utf-8')))
    # exact best family only; remove top 1/3 winners from event sequence, recompute from 2m
    z=[r for r in rows if r['side']=='DOWN' and float(r['threshold'])==.03 and int(r['hold'])==3 and int(r['topN'])==5]
    rets=[float(r['event_ret_net']) for r in z]
    def eq(rs):
        c=START_CAP;pk=c;dd=0
        for x in rs:c*=1+x;pk=max(pk,c);dd=min(dd,c/pk-1)
        return c,c/START_CAP-1,dd
    s=sorted(rets,reverse=True)
    return {'n':len(rets),'base':eq(rets),'remove_best1':eq([x for x in rets if x!=s[0]]) if rets else None,'remove_best3':eq([x for i,x in enumerate(rets) if i not in {rets.index(v) for v in s[:3]}]) if len(rets)>=3 else None}

def main():
    panel,qa=fetch_twse(); sox=fetch_sox()
    errs=sum(s=='ERROR' for _,s,_,_ in qa); passes=sum(s=='PASS' for _,s,_,_ in qa)
    grid=[]
    for cost in (.004,.006,.008,.01):
      for th in (.015,.02,.03):
       for hold in (2,3,4,5):
        for topN in (3,5,10):
         r=backtest(panel,sox,cost,hold,th,topN);grid.append({'cost':cost,'th':th,'hold':hold,'topN':topN,**{k:v for k,v in r.items() if k!='events'}})
    # fixed 2026 champion on untouched 24-25 plus by-year equity path derived from events
    fixed=backtest(panel,sox,.004,3,.03,5)
    events=fixed['events']; yearly={}
    for y in ('2024','2025'):
        c=START_CAP;pk=c;dd=0;ys=[x for x in events if x['entry'].startswith(y)]
        for e in ys:c*=1+e['ret'];pk=max(pk,c);dd=min(dd,c/pk-1)
        yearly[y]={'final':c,'ret':c/START_CAP-1,'mdd':dd,'n':len(ys),'wins':sum(e['ret']>0 for e in ys)}
    # neighborhood robustness summary
    positive=sum(r['ret']>0 for r in grid); gt20=sum(r['ret']>.2 for r in grid)
    out={'qa':{'pass_days':passes,'errors':errs,'panel_rows':len(panel),'sox_rows':len(sox)},'fixed_2026_champion_on_2024_2025':{k:v for k,v in fixed.items() if k!='events'},'yearly_reset_2m':yearly,'grid':{'tested':len(grid),'positive':positive,'gt20pct':gt20,'top10':sorted(grid,key=lambda x:x['ret'],reverse=True)[:10]},'stress_2026':stress_2026()}
    (OUT/'summary.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
    with (OUT/'grid.csv').open('w',newline='',encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=['cost','th','hold','topN','final','ret','mdd','wins','n']);w.writeheader();w.writerows(grid)
    with (OUT/'fixed_events.csv').open('w',newline='',encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=['entry','exit','ret','cap','shock','names']);w.writeheader();w.writerows(events)
    print(json.dumps(out,ensure_ascii=False,indent=2))
    if errs or passes<470 or len(sox)<480:raise SystemExit(2)

if __name__=='__main__':main()
